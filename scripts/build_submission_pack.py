#!/usr/bin/env python3
"""Build Day 5 submission PDFs + case-study PPTX from markdown sources.

Dependencies (install if missing):
  pip install fpdf2 markdown python-pptx

Usage (from repo root):
  python scripts/build_submission_pack.py
"""

from __future__ import annotations

import re
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "docs" / "submission"

try:
    from fpdf import FPDF
except ImportError:
    print("Install fpdf2: pip install fpdf2", file=sys.stderr)
    sys.exit(1)

try:
    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.util import Inches, Pt
except ImportError:
    print("Install python-pptx: pip install python-pptx", file=sys.stderr)
    sys.exit(1)


def read_md(*rel_paths: str) -> str:
    parts: list[str] = []
    for rel in rel_paths:
        path = ROOT / rel
        if not path.exists():
            raise FileNotFoundError(path)
        parts.append(f"# Source: {rel}\n\n{path.read_text(encoding='utf-8')}")
    return "\n\n---\n\n".join(parts)


def strip_md(text: str) -> list[tuple[str, str]]:
    """Return list of (style, line) where style in h1|h2|h3|body|bullet|table|hr|code."""
    lines_out: list[tuple[str, str]] = []
    in_code = False
    for raw in text.splitlines():
        line = raw.rstrip()
        if line.startswith("```"):
            in_code = not in_code
            continue
        if in_code:
            lines_out.append(("code", line))
            continue
        if line.strip() == "---":
            lines_out.append(("hr", ""))
            continue
        if line.startswith("# "):
            lines_out.append(("h1", line[2:].strip()))
            continue
        if line.startswith("## "):
            lines_out.append(("h2", line[3:].strip()))
            continue
        if line.startswith("### "):
            lines_out.append(("h3", line[4:].strip()))
            continue
        if re.match(r"^[-*] ", line):
            lines_out.append(("bullet", re.sub(r"^[-*] ", "", line)))
            continue
        if re.match(r"^- \[[ xX]\] ", line):
            checked = "[x]" in line[:6].lower()
            body = re.sub(r"^- \[[ xX]\] ", "", line)
            lines_out.append(("bullet", f"{'☑' if checked else '☐'} {body}"))
            continue
        if "|" in line and line.strip().startswith("|"):
            cells = [c.strip() for c in line.strip().strip("|").split("|")]
            if all(re.match(r"^:?-+:?$", c) for c in cells):
                continue
            lines_out.append(("table", "  |  ".join(cells)))
            continue
        # inline cleanup
        cleaned = re.sub(r"\*\*(.+?)\*\*", r"\1", line)
        cleaned = re.sub(r"`([^`]+)`", r"\1", cleaned)
        cleaned = re.sub(r"\[([^\]]+)\]\([^)]+\)", r"\1", cleaned)
        if cleaned.strip():
            lines_out.append(("body", cleaned))
        else:
            lines_out.append(("body", ""))
    return lines_out


class PackPDF(FPDF):
    def footer(self) -> None:
        self.set_y(-12)
        self.set_font("Helvetica", size=8)
        self.set_text_color(100, 100, 100)
        self.cell(0, 8, f"AccountFlow OS - page {self.page_no()}", align="C")


def md_to_pdf(title: str, markdown: str, dest: Path) -> None:
    pdf = PackPDF(format="A4")
    pdf.set_auto_page_break(auto=True, margin=18)
    pdf.set_margins(left=14, top=14, right=14)
    pdf.add_page()
    pdf.set_font("Helvetica", "B", 16)
    pdf.multi_cell(0, 9, _safe(title))
    pdf.ln(4)

    usable = pdf.epw

    for style, text in strip_md(markdown):
        text = _safe(text)
        if style == "hr":
            pdf.ln(2)
            y = pdf.get_y()
            pdf.set_draw_color(180, 180, 180)
            pdf.line(pdf.l_margin, y, pdf.l_margin + usable, y)
            pdf.ln(4)
            continue
        # Avoid zero-width multi_cell when x drifts past right margin
        if pdf.get_x() > pdf.l_margin + 1:
            pdf.ln()
        pdf.set_x(pdf.l_margin)
        if style == "h1":
            pdf.ln(3)
            pdf.set_font("Helvetica", "B", 14)
            pdf.multi_cell(usable, 8, text)
            pdf.ln(1)
        elif style == "h2":
            pdf.ln(2)
            pdf.set_font("Helvetica", "B", 12)
            pdf.multi_cell(usable, 7, text)
            pdf.ln(1)
        elif style == "h3":
            pdf.set_font("Helvetica", "B", 11)
            pdf.multi_cell(usable, 6, text)
        elif style == "bullet":
            pdf.set_font("Helvetica", size=10)
            pdf.multi_cell(usable, 5, f"  * {text}")
        elif style == "table":
            pdf.set_font("Courier", size=7)
            # Soft-wrap long table rows
            while len(text) > 110:
                pdf.multi_cell(usable, 4, text[:110])
                text = text[110:]
            pdf.multi_cell(usable, 4, text or " ")
        elif style == "code":
            pdf.set_font("Courier", size=8)
            chunk = text or " "
            while len(chunk) > 95:
                pdf.multi_cell(usable, 4.5, chunk[:95])
                chunk = chunk[95:]
            pdf.multi_cell(usable, 4.5, chunk)
        else:
            pdf.set_font("Helvetica", size=10)
            if text:
                pdf.multi_cell(usable, 5, text)
            else:
                pdf.ln(2)

    dest.parent.mkdir(parents=True, exist_ok=True)
    pdf.output(str(dest))
    print(f"Wrote {dest.relative_to(ROOT)}")


def _safe(s: str) -> str:
    # Helvetica core fonts are Latin-1; replace common Unicode
    replacements = {
        "\u2192": "->",
        "\u2190": "<-",
        "\u2014": "-",
        "\u2013": "-",
        "\u2026": "...",
        "\u201c": '"',
        "\u201d": '"',
        "\u2018": "'",
        "\u2019": "'",
        "\u2022": "*",
        "\u2611": "[x]",
        "\u2610": "[ ]",
        "\u00d7": "x",
        "\u2265": ">=",
        "\u2264": "<=",
        "\u2248": "~",
        "\u00b7": "-",
        "\u00a0": " ",
        "\ufeff": "",
    }
    for a, b in replacements.items():
        s = s.replace(a, b)
    return s.encode("latin-1", errors="replace").decode("latin-1")


def build_pptx(dest: Path) -> None:
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    slides_spec = [
        (
            "AccountFlow OS",
            "Post-call OS: transcript + HubSpot + SOW -> Scope Verifier -> approve -> Gmail / CRM / Jira\n"
            "Live: accountflow-web.vercel.app  |  API: accountflow-api.vercel.app",
        ),
        (
            "Problem & baseline",
            "Alex (AE/PM): 8-12 client calls/week, 30-40 min admin each.\n"
            "Measured means (BASELINE.md): Manual 38 min | ChatGPT 22 min | still copy-paste, missed scope.",
        ),
        (
            "Architecture (v1)",
            "Ingest: record/upload + STT, HubSpot deal, SOW\n"
            "Verify: Scope flags (IN/OUT/TIMELINE/BUDGET) + hallucination grader\n"
            "Approve: human required — no auto-send\n"
            "Execute: Gmail, HubSpot write, Jira create\n"
            "Stack: LangGraph + FastAPI + Next.js + Turso (prod) on Vercel",
        ),
        (
            "Live flow",
            "1. New Run — select/create HubSpot deal + paste demo transcript/SOW\n"
            "2. Scope Verifier flags out-of-scope / timeline issues\n"
            "3. Approval Inbox — edit CRM stage dropdowns, email, tasks\n"
            "4. Execute — nothing runs until Approve\n"
            "Fixtures: samples/demo/",
        ),
        (
            "Evaluation",
            "Automated suite: 12/12 PASS (eval/run_eval.py)\n"
            "Failures fixed: TC-06 hallucinated deadline, TC-05 owner ambiguity, TC-09 Gmail retries\n"
            "Not measured in v1: per-call $, production P95 latency",
        ),
        (
            "Results",
            "Primary metric (mean): 38 min -> 6 min (mock path)\n"
            "Target <8 min — met on baseline transcripts\n"
            "CRM quality 4.5/5 | Scope caught on planted cases | Human approve ~1 min overhead",
        ),
        (
            "Limitation & next 2 weeks",
            "Limitation: Scope Verifier tracks SOW quality; human gate stays.\n"
            "Week 1: 5 live proxy calls, cost logging, tune false positives\n"
            "Week 2: optional recording ingest / Slack notify, commitment tracker, HubSpot OAuth review",
        ),
    ]

    for title, body in slides_spec:
        slide = prs.slides.add_slide(prs.slide_layouts[6])  # blank
        # atmosphere bar
        shape = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE,
            Inches(0),
            Inches(0),
            Inches(13.333),
            Inches(1.15),
        )
        shape.fill.solid()
        shape.fill.fore_color.rgb = RGBColor(0x1A, 0x3A, 0x4A)
        shape.line.fill.background()

        title_box = slide.shapes.add_textbox(Inches(0.6), Inches(0.28), Inches(12), Inches(0.7))
        tf = title_box.text_frame
        p = tf.paragraphs[0]
        p.text = title
        p.font.size = Pt(28)
        p.font.bold = True
        p.font.color.rgb = RGBColor(0xF5, 0xF0, 0xE8)
        p.font.name = "Calibri"

        body_box = slide.shapes.add_textbox(Inches(0.7), Inches(1.6), Inches(11.8), Inches(5.2))
        btf = body_box.text_frame
        btf.word_wrap = True
        first = True
        for para_text in body.split("\n"):
            para = btf.paragraphs[0] if first else btf.add_paragraph()
            first = False
            para.text = para_text
            para.font.size = Pt(18)
            para.font.color.rgb = RGBColor(0x1A, 0x2A, 0x33)
            para.font.name = "Calibri"
            para.space_after = Pt(10)

    dest.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(dest))
    print(f"Wrote {dest.relative_to(ROOT)}")


def main() -> None:
    OUT.mkdir(parents=True, exist_ok=True)

    packs = [
        (
            "01_Working_System.pdf",
            "01 Working System — AccountFlow OS",
            [
                "README.md",
                "docs/RUNBOOK.md",
                "docs/DECISIONS.md",
                "samples/README.md",
            ],
        ),
        (
            "02_Evaluation_Package.pdf",
            "02 Evaluation Package — AccountFlow OS",
            [
                "docs/evaluation/BASELINE.md",
                "docs/evaluation/RUBRIC.md",
                "docs/evaluation/METRICS.md",
                "docs/evaluation/RESULTS.md",
                "docs/evaluation/FAILURES.md",
                "docs/evaluation/RESULTS_TABLE.md",
            ],
        ),
        (
            "03_Case_Study.pdf",
            "03 Case Study — AccountFlow OS",
            ["docs/case-study/CASE_STUDY.md"],
        ),
        (
            "04_AI_Collaboration_Note.pdf",
            "04 AI Collaboration Note",
            ["docs/case-study/AI_COLLABORATION.md"],
        ),
        (
            "05_Demo_Script.pdf",
            "05 Demo Script (for recording)",
            ["docs/case-study/DEMO_SCRIPT.md"],
        ),
    ]

    for filename, title, sources in packs:
        # skip missing optional table
        existing = [s for s in sources if (ROOT / s).exists()]
        md_to_pdf(title, read_md(*existing), OUT / filename)

    build_pptx(OUT / "AccountFlow_OS_Case_Study.pptx")

    readme = OUT / "README.md"
    readme.write_text(
        "# Submission pack\n\n"
        "Regenerate with:\n\n"
        "```bash\n"
        "pip install fpdf2 python-pptx markdown\n"
        "python scripts/build_submission_pack.py\n"
        "```\n\n"
        "| File | Content |\n"
        "|------|--------|\n"
        "| `01_Working_System.pdf` | README, RUNBOOK, DECISIONS, samples |\n"
        "| `02_Evaluation_Package.pdf` | Baseline, rubric, metrics, results, failures |\n"
        "| `03_Case_Study.pdf` | Case study |\n"
        "| `04_AI_Collaboration_Note.pdf` | AI collaboration note |\n"
        "| `05_Demo_Script.pdf` | Demo recording script |\n"
        "| `AccountFlow_OS_Case_Study.pptx` | ~7-slide skim deck |\n",
        encoding="utf-8",
    )
    print(f"Wrote {readme.relative_to(ROOT)}")
    print("Done.")


if __name__ == "__main__":
    main()
